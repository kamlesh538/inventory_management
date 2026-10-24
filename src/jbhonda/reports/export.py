"""Data export to Excel (one workbook, a sheet per report) or CSV (a file per report).

Amounts are written as numbers in rupees, so they add up in Excel; dates as real dates.
"""

import csv
import sqlite3
from collections.abc import Callable
from dataclasses import dataclass
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

MONEY, DATE, TEXT, NUMBER = "money", "date", "text", "number"


@dataclass(frozen=True)
class Report:
    title: str
    columns: list[tuple[str, str]]  # (heading, kind)
    rows: list[tuple]


def _range(args: list, column: str, date_from: date | None, date_to: date | None) -> str:
    where = []
    if date_from:
        where.append(f"{column} >= ?")
        args.append(date_from.isoformat())
    if date_to:
        where.append(f"{column} <= ?")
        args.append(date_to.isoformat())
    return " AND ".join(where) or "1"


def sales_register(conn: sqlite3.Connection, date_from: date | None, date_to: date | None) -> Report:
    args: list = []
    rows = conn.execute(
        f"""
        SELECT i.invoice_no, i.invoice_date, CASE i.invoice_type WHEN 'VEHICLE' THEN 'Vehicle' WHEN 'SERVICE' THEN 'Service' ELSE 'Parts' END,
               i.bill_name, i.bill_phone, i.bill_gstin, i.total_taxable, i.total_cgst, i.total_sgst, i.round_off,
               i.grand_total,
               (SELECT coalesce(sum(amount), 0) FROM payment p WHERE p.invoice_id = i.id),
               CASE i.status WHEN 'ACTIVE' THEN 'Active' ELSE 'Cancelled' END, i.hypothecated_to, i.cancel_reason
        FROM invoice i WHERE {_range(args, "i.invoice_date", date_from, date_to)}
        ORDER BY i.invoice_date, i.invoice_no
        """, args).fetchall()
    return Report("Sales register", [
        ("Invoice no", TEXT), ("Date", DATE), ("Type", TEXT), ("Customer", TEXT), ("Mobile", TEXT),
        ("Customer GSTIN", TEXT), ("Taxable value", MONEY), ("CGST", MONEY), ("SGST", MONEY), ("Round off", MONEY),
        ("Invoice total", MONEY), ("Paid", MONEY), ("Status", TEXT), ("Hypothecated to", TEXT),
        ("Cancel reason", TEXT),
    ], [tuple(r) for r in rows])


def gst_summary(conn: sqlite3.Connection, date_from: date | None, date_to: date | None) -> Report:
    """HSN-wise outward supplies for GSTR-1 (active invoices only)."""
    args: list = []
    rows = conn.execute(
        f"""
        SELECT ii.hsn_code, ii.gst_rate_bp / 100.0, sum(ii.quantity), sum(ii.taxable_value), sum(ii.cgst),
               sum(ii.sgst), sum(ii.cgst + ii.sgst), sum(ii.line_total)
        FROM invoice_item ii JOIN invoice i ON i.id = ii.invoice_id
        WHERE i.status = 'ACTIVE' AND {_range(args, "i.invoice_date", date_from, date_to)}
        GROUP BY ii.hsn_code, ii.gst_rate_bp ORDER BY ii.hsn_code, ii.gst_rate_bp
        """, args).fetchall()
    return Report("GST HSN summary", [
        ("HSN", TEXT), ("GST rate %", NUMBER), ("Quantity", NUMBER), ("Taxable value", MONEY), ("CGST", MONEY),
        ("SGST", MONEY), ("Total tax", MONEY), ("Total value", MONEY),
    ], [tuple(r) for r in rows])


def payments_report(conn: sqlite3.Connection, date_from: date | None, date_to: date | None) -> Report:
    args: list = []
    rows = conn.execute(
        f"""
        SELECT p.receipt_no, p.payment_date, coalesce(i.invoice_no, ''), coalesce(b.booking_no, ''),
               coalesce(c.name, i.bill_name), replace(lower(p.payment_type), '_', ' '),
               replace(lower(p.payment_mode), '_', ' '), p.reference_no, p.amount
        FROM payment p
        LEFT JOIN invoice i ON i.id = p.invoice_id
        LEFT JOIN booking b ON b.id = p.booking_id
        LEFT JOIN customer c ON c.id = p.customer_id
        WHERE {_range(args, "p.payment_date", date_from, date_to)}
        ORDER BY p.payment_date, p.id
        """, args).fetchall()
    return Report("Payments", [
        ("Receipt no", TEXT), ("Date", DATE), ("Invoice no", TEXT), ("Booking no", TEXT), ("Customer", TEXT),
        ("Type", TEXT), ("Mode", TEXT), ("Reference", TEXT), ("Amount", MONEY),
    ], [tuple(r) for r in rows])


def vehicle_stock(conn: sqlite3.Connection, date_from: date | None = None, date_to: date | None = None) -> Report:
    rows = conn.execute(
        """
        SELECT m.model_name, m.variant, v.color, v.chassis_number, v.engine_number, v.key_number,
               v.manufacturing_month, replace(lower(v.status), '_', ' '), v.purchase_invoice_no, v.purchase_date,
               v.purchase_taxable_value, v.purchase_gst_amount, s.invoice_no, s.invoice_date, s.bill_name,
               v.registration_number
        FROM vehicle v JOIN vehicle_model m ON m.id = v.vehicle_model_id
        LEFT JOIN (SELECT ii.vehicle_id, i.invoice_no, i.invoice_date, i.bill_name FROM invoice_item ii
                   JOIN invoice i ON i.id = ii.invoice_id WHERE i.status = 'ACTIVE') s ON s.vehicle_id = v.id
        ORDER BY v.status, m.model_name, v.color, v.chassis_number
        """).fetchall()
    return Report("Vehicle stock", [
        ("Model", TEXT), ("Variant", TEXT), ("Colour", TEXT), ("Chassis no", TEXT), ("Engine no", TEXT),
        ("Key no", TEXT), ("Mfg month", TEXT), ("Status", TEXT), ("Purchase invoice", TEXT), ("Purchase date", DATE),
        ("Purchase price", MONEY), ("Purchase GST", MONEY), ("Sale invoice", TEXT), ("Sale date", DATE),
        ("Customer", TEXT), ("Registration no", TEXT),
    ], [tuple(r) for r in rows])


def parts_stock(conn: sqlite3.Connection, date_from: date | None = None, date_to: date | None = None) -> Report:
    rows = conn.execute(
        """
        SELECT part_number, name, lower(category), hsn_code, gst_rate_bp / 100.0, unit, stock_qty, reorder_level,
               rack_location, mrp, sale_price, avg_cost, stock_qty * avg_cost,
               CASE WHEN stock_qty <= reorder_level THEN 'Yes' ELSE '' END, CASE is_active WHEN 1 THEN '' ELSE 'Yes' END
        FROM part ORDER BY name COLLATE NOCASE
        """).fetchall()
    return Report("Parts stock", [
        ("Part no", TEXT), ("Name", TEXT), ("Category", TEXT), ("HSN", TEXT), ("GST rate %", NUMBER), ("Unit", TEXT),
        ("Stock", NUMBER), ("Reorder level", NUMBER), ("Rack", TEXT), ("MRP", MONEY), ("Sale price", MONEY),
        ("Average cost", MONEY), ("Stock value", MONEY), ("Low stock", TEXT), ("Inactive", TEXT),
    ], [tuple(r) for r in rows])


def customers_report(conn: sqlite3.Connection, date_from: date | None = None, date_to: date | None = None) -> Report:
    rows = conn.execute(
        """
        SELECT name, relation_type, relation_name, phone, alt_phone, email, address_line1, address_line2, city,
               district, state, pincode, gstin, id_proof_type, id_proof_number, substr(created_at, 1, 10)
        FROM customer ORDER BY name COLLATE NOCASE
        """).fetchall()
    return Report("Customers", [
        ("Name", TEXT), ("Relation", TEXT), ("Relation name", TEXT), ("Mobile", TEXT), ("Alt. phone", TEXT),
        ("Email", TEXT), ("Address 1", TEXT), ("Address 2", TEXT), ("City", TEXT), ("District", TEXT),
        ("State", TEXT), ("PIN", TEXT), ("GSTIN", TEXT), ("ID proof", TEXT), ("ID number", TEXT), ("Added on", DATE),
    ], [tuple(r) for r in rows])


def bookings_report(conn: sqlite3.Connection, date_from: date | None, date_to: date | None) -> Report:
    args: list = []
    rows = conn.execute(
        f"""
        SELECT b.booking_no, b.booking_date, c.name, c.phone, m.model_name || ' ' || m.variant, b.preferred_color,
               v.chassis_number, b.quoted_price,
               (SELECT coalesce(sum(amount), 0) FROM payment p WHERE p.booking_id = b.id
                  AND p.payment_type = 'BOOKING_ADVANCE'),
               lower(b.status), i.invoice_no, b.cancel_reason
        FROM booking b JOIN customer c ON c.id = b.customer_id JOIN vehicle_model m ON m.id = b.vehicle_model_id
        LEFT JOIN vehicle v ON v.id = b.vehicle_id LEFT JOIN invoice i ON i.id = b.invoice_id
        WHERE {_range(args, "b.booking_date", date_from, date_to)}
        ORDER BY b.booking_date, b.booking_no
        """, args).fetchall()
    return Report("Bookings", [
        ("Booking no", TEXT), ("Date", DATE), ("Customer", TEXT), ("Mobile", TEXT), ("Model", TEXT),
        ("Colour wanted", TEXT), ("Vehicle", TEXT), ("Quoted price", MONEY), ("Advance paid", MONEY),
        ("Status", TEXT), ("Invoice no", TEXT), ("Cancel reason", TEXT),
    ], [tuple(r) for r in rows])


# Order shown on screen. Stock and customer reports ignore the date range (they're "as of now").
REPORTS: dict[str, tuple[str, Callable[..., Report], bool]] = {
    "sales": ("Sales register", sales_register, True),
    "gst": ("GST HSN summary", gst_summary, True),
    "payments": ("Payments received", payments_report, True),
    "bookings": ("Bookings", bookings_report, True),
    "vehicles": ("Vehicle stock (current)", vehicle_stock, False),
    "parts": ("Parts stock (current)", parts_stock, False),
    "customers": ("Customers", customers_report, False),
}


def build_reports(conn: sqlite3.Connection, keys: list[str], date_from: date | None, date_to: date | None) -> list[Report]:
    return [REPORTS[k][1](conn, date_from, date_to) for k in keys]


def _cell(value, kind: str):
    if value is None or value == "":
        return None
    if kind == MONEY:
        return value / 100
    if kind == DATE:
        return date.fromisoformat(str(value)[:10])
    return value


def write_xlsx(path: Path, reports: list[Report], subtitle: str = "") -> Path:
    wb = Workbook()
    wb.remove(wb.active)
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F2933")
    for report in reports:
        ws = wb.create_sheet(report.title[:31])
        ws.append([heading for heading, _ in report.columns])
        for cell in ws[1]:
            cell.font, cell.fill = header_font, header_fill
        for row in report.rows:
            ws.append([_cell(v, kind) for v, (_, kind) in zip(row, report.columns)])
        for i, (heading, kind) in enumerate(report.columns, start=1):
            letter = get_column_letter(i)
            fmt = {"money": "#,##0.00", "date": "DD-MM-YYYY"}.get(kind)
            if fmt:
                for cell in ws[letter][1:]:
                    cell.number_format = fmt
            width = max([len(heading)] + [len(str(r[i - 1])) for r in report.rows[:500] if r[i - 1] is not None])
            ws.column_dimensions[letter].width = min(max(width + 2, 10), 45)
        money_cols = [i for i, (_, kind) in enumerate(report.columns, start=1) if kind == MONEY]
        if report.rows and money_cols:
            total_row = ws.max_row + 1
            ws.cell(total_row, 1, "Total").font = Font(bold=True)
            for i in money_cols:
                letter = get_column_letter(i)
                cell = ws.cell(total_row, i, f"=SUM({letter}2:{letter}{total_row - 1})")
                cell.font, cell.number_format = Font(bold=True), "#,##0.00"
        ws.freeze_panes = "A2"
        if subtitle:
            ws.oddHeader.center.text = f"{report.title} - {subtitle}"
    wb.save(path)
    return path


def write_csv(folder: Path, reports: list[Report], stem: str) -> list[Path]:
    folder.mkdir(parents=True, exist_ok=True)
    written = []
    for report in reports:
        path = folder / f"{stem} - {report.title}.csv"
        # utf-8-sig so Excel shows ₹ and Hindi names correctly when the CSV is opened directly
        with path.open("w", newline="", encoding="utf-8-sig") as f:
            writer = csv.writer(f)
            writer.writerow([heading for heading, _ in report.columns])
            for row in report.rows:
                values = [_cell(v, kind) for v, (_, kind) in zip(row, report.columns)]
                writer.writerow(["" if v is None else v.strftime("%d-%m-%Y") if isinstance(v, date) else
                                 f"{v:.2f}" if kind == MONEY else v
                                 for v, (_, kind) in zip(values, report.columns)])
        written.append(path)
    return written
