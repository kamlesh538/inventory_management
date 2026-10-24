import csv
import sqlite3
from datetime import date, datetime, timedelta

import pytest
from openpyxl import load_workbook

from jbhonda import config
from jbhonda.core.backup import KEEP_AUTO, auto_backup, backup_now, check_backup_file, list_backups, restore
from jbhonda.core.customers import save_customer
from jbhonda.core.errors import BusinessError
from jbhonda.core.inventory import save_part, save_vehicle, save_vehicle_model
from jbhonda.core.sales import PartLine, PaymentIn, cancel_invoice, complete_part_sale
from jbhonda.core.vehicle_sale import complete_vehicle_sale
from jbhonda.reports.export import build_reports, write_csv, write_xlsx
from jbhonda.ui import data_page, widgets
from jbhonda.ui.data_page import DataPage

DAY = date(2026, 10, 10)


@pytest.fixture
def shop(conn):
    m = save_vehicle_model(conn, model_name="Activa", category="SCOOTER", gst_rate_bp=1800, ex_showroom_price=8500000)
    v = save_vehicle(conn, vehicle_model_id=m, chassis_number="CH1", engine_number="EN1", color="Red",
                     purchase_taxable_value=6800000)
    save_vehicle(conn, vehicle_model_id=m, chassis_number="CH2", engine_number="EN2", color="Blue")
    c = save_customer(conn, name="Ramesh", phone="9876543210")
    oil = save_part(conn, part_number="OIL", name="Oil", hsn_code="2710", gst_rate_bp=1800, sale_price=40000,
                    opening_qty=5, opening_unit_cost=30000)
    complete_vehicle_sale(conn, customer_id=c, vehicle_id=v, payments=[PaymentIn("CASH", 8500000)], on=DAY)
    p = complete_part_sale(conn, lines=[PartLine(oil, 2)], payments=[PaymentIn("CASH", 80000)], on=DAY)
    cancelled = complete_part_sale(conn, lines=[PartLine(oil, 1)], payments=[PaymentIn("CASH", 40000)], on=DAY)
    cancel_invoice(conn, cancelled, "mistake")
    complete_part_sale(conn, lines=[PartLine(oil, 1)], payments=[PaymentIn("CASH", 40000)], on=DAY + timedelta(days=30))
    return {"part_invoice": p}


# ---------------------------------------------------------------- export

def test_reports_respect_period_and_cancellations(conn, shop):
    sales, gst, vehicles = build_reports(conn, ["sales", "gst", "vehicles"], date(2026, 10, 1), date(2026, 10, 31))
    assert len(sales.rows) == 3  # two parts bills (one cancelled) + vehicle; November bill excluded
    assert [r[12] for r in sales.rows].count("Cancelled") == 1
    by_hsn = {r[0]: r for r in gst.rows}
    assert set(by_hsn) == {"2710", "8711"}  # cancelled bill not counted
    assert by_hsn["2710"][2] == 2 and by_hsn["2710"][7] == 80000
    assert len(vehicles.rows) == 2  # stock is "as of now", not filtered by period


def test_excel_export(conn, shop, tmp_path):
    path = write_xlsx(tmp_path / "x.xlsx", build_reports(conn, ["sales", "parts"], None, None), subtitle="all")
    wb = load_workbook(path)
    assert wb.sheetnames == ["Sales register", "Parts stock"]
    ws = wb["Sales register"]
    assert ws["A1"].value == "Invoice no" and ws.freeze_panes == "A2"
    assert isinstance(ws["B2"].value, datetime)  # real dates
    totals = [ws.cell(r, 11).value for r in range(2, ws.max_row)]
    assert 85000 in totals  # rupees, not paise
    assert ws.cell(ws.max_row, 1).value == "Total"
    assert ws.cell(ws.max_row, 11).value.startswith("=SUM(K2:K")


def test_csv_export(conn, shop, tmp_path):
    (path,) = write_csv(tmp_path, build_reports(conn, ["payments"], None, None), "JBH")
    with path.open(encoding="utf-8-sig") as f:
        rows = list(csv.reader(f))
    assert rows[0][:3] == ["Receipt no", "Date", "Invoice no"]
    assert rows[1][1] == "10-10-2026" and rows[1][-1] == "85000.00"


# ---------------------------------------------------------------- backup

def test_backup_now_to_any_folder(conn, shop, tmp_path):
    path = backup_now(conn, tmp_path / "pendrive")
    assert check_backup_file(path) >= 1
    copy = sqlite3.connect(path)
    assert copy.execute("SELECT count(*) FROM invoice").fetchone()[0] == 4
    copy.close()


def test_auto_backup_once_a_day_and_prunes(conn, shop):
    assert auto_backup(conn) is not None
    assert auto_backup(conn) is None  # already done today
    folder = config.backup_dir()
    for i in range(KEEP_AUTO + 5):
        stamp = (datetime(2025, 1, 1) + timedelta(days=i)).strftime("%Y%m%d-%H%M%S")
        (folder / f"data-{stamp}-auto.db").write_bytes(b"x")
    (folder / "data-20200101-000000-manual.db").write_bytes(b"x")
    auto_backup(conn)
    autos = [b for b in list_backups() if b.automatic]
    assert len(autos) == KEEP_AUTO
    assert autos[0].taken_at.date() == date.today()  # newest kept
    assert any(not b.automatic for b in list_backups())  # manual backups never pruned


def test_restore_replaces_data_and_keeps_safety_copy(conn, shop, tmp_path):
    backup = backup_now(conn, tmp_path)
    save_customer(conn, name="Added After Backup", phone="9123456789")
    safety = restore(conn, backup)
    assert conn.execute("SELECT count(*) FROM customer WHERE name = 'Added After Backup'").fetchone()[0] == 0
    old = sqlite3.connect(safety)
    assert old.execute("SELECT count(*) FROM customer WHERE name = 'Added After Backup'").fetchone()[0] == 1
    old.close()
    assert conn.execute("PRAGMA foreign_keys").fetchone()[0] == 1


def test_restore_rejects_wrong_files(conn, tmp_path):
    junk = tmp_path / "junk.db"
    junk.write_bytes(b"not a database")
    with pytest.raises(BusinessError, match="not a backup"):
        restore(conn, junk)
    other = tmp_path / "other.db"
    sqlite3.connect(other).execute("CREATE TABLE x (a)").connection.close()
    with pytest.raises(BusinessError, match="not a backup"):
        restore(conn, other)
    newer = backup_now(conn, tmp_path / "n")
    c = sqlite3.connect(newer)
    c.execute("PRAGMA user_version = 999")
    c.close()
    with pytest.raises(BusinessError, match="newer version"):
        restore(conn, newer)


# ---------------------------------------------------------------- screen

@pytest.fixture
def messages(monkeypatch):
    shown = []
    monkeypatch.setattr(widgets, "show_error", lambda parent, msg: shown.append(("error", msg)))
    monkeypatch.setattr(widgets, "show_info", lambda parent, msg: shown.append(("info", msg)))
    monkeypatch.setattr(widgets, "confirm", lambda parent, msg: shown.append(("confirm", msg)) or True)
    return shown


def test_data_page_export_excel(qtbot, conn, shop, messages, opened_files, tmp_path, monkeypatch):
    target = tmp_path / "out.xlsx"
    monkeypatch.setattr(data_page, "_choose_save_path", lambda parent, suggested: target)
    page = DataPage(conn)
    qtbot.addWidget(page)
    page.period.setValue("ALL")
    assert page.export() == [target]
    assert load_workbook(target).sheetnames == ["Sales register", "GST HSN summary"]
    assert opened_files == [target]


def test_data_page_requires_a_report(qtbot, conn, messages):
    page = DataPage(conn)
    qtbot.addWidget(page)
    for box in page.report_boxes.values():
        box.setChecked(False)
    assert page.export() == []
    assert messages[-1] == ("error", "Tick at least one report to export.")


def test_data_page_backup_and_restore(qtbot, conn, shop, messages):
    page = DataPage(conn)
    qtbot.addWidget(page)
    path = page.make_backup(None)
    assert page.backups.rowCount() == 1
    save_customer(conn, name="Later", phone="9123456780")
    assert page.restore_from_file(path)
    assert conn.execute("SELECT count(*) FROM customer WHERE name = 'Later'").fetchone()[0] == 0
    assert messages[-1][0] == "info" and "restored" in messages[-1][1]
