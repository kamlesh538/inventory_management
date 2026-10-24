from datetime import date

import pytest

from jbhonda.core.customers import save_customer
from jbhonda.core.inventory import save_part, save_vehicle, save_vehicle_model
from jbhonda.core.masters import save_financer
from jbhonda.core.sales import PartLine, PaymentIn, cancel_invoice, complete_part_sale
from jbhonda.core.sales_report import Filters, sales_report
from jbhonda.core.service import add_job_labour, bill_job, create_job
from jbhonda.core.vehicle_sale import LoanIn, complete_vehicle_sale

OCT = date(2026, 10, 10)


@pytest.fixture
def sales(conn):
    m = save_vehicle_model(conn, model_name="Activa", category="SCOOTER", gst_rate_bp=1800, ex_showroom_price=8500000)
    v1 = save_vehicle(conn, vehicle_model_id=m, chassis_number="C1", engine_number="E1", color="Red")
    v2 = save_vehicle(conn, vehicle_model_id=m, chassis_number="C2", engine_number="E2", color="Red")
    c = save_customer(conn, name="Ram", phone="9876543210")
    f = save_financer(conn, name="HDFC")
    oil = save_part(conn, part_number="OIL", name="Oil", hsn_code="2710", gst_rate_bp=1800, sale_price=40000,
                    opening_qty=10)
    # financed vehicle: 15,000 cash + 70,000 loan (received)
    complete_vehicle_sale(conn, customer_id=c, vehicle_id=v1, payments=[PaymentIn("CASH", 1500000)],
                          loan=LoanIn(financer_id=f, loan_amount=7000000, disbursed=True), on=OCT)
    # direct vehicle: split cash + UPI
    complete_vehicle_sale(conn, customer_id=c, vehicle_id=v2,
                          payments=[PaymentIn("CASH", 5000000), PaymentIn("UPI", 3500000, "U1")], on=OCT)
    # parts: one UPI bill, one cancelled bill, one next month
    complete_part_sale(conn, lines=[PartLine(oil, 1)], payments=[PaymentIn("UPI", 40000, "U2")], on=OCT)
    dead = complete_part_sale(conn, lines=[PartLine(oil, 1)], payments=[PaymentIn("CASH", 40000)], on=OCT)
    cancel_invoice(conn, dead, "mistake")
    complete_part_sale(conn, lines=[PartLine(oil, 1)], payments=[PaymentIn("CASH", 40000)], on=date(2026, 11, 2))
    # service on credit
    job = create_job(conn, customer_id=c, vehicle_description="Activa")
    add_job_labour(conn, job, description="General service", amount=35000)
    bill_job(conn, job, on=OCT)
    return conn


def _report(conn, **kw):
    return sales_report(conn, Filters(date_from=date(2026, 10, 1), date_to=date(2026, 10, 31), **kw))


def test_totals_for_the_month(sales):
    r = _report(sales)
    assert r.count == 4  # cancelled and November bills excluded
    assert r.sale_value == 8500000 * 2 + 40000 + 35000
    assert r.received == 8500000 * 2 + 40000
    assert r.outstanding == 35000  # the unpaid service bill
    assert r.taxable + r.gst == r.sale_value
    assert {k: v["count"] for k, v in r.by_type.items()} == {"VEHICLE": 2, "PART": 1, "SERVICE": 1}
    assert r.by_mode == {"CASH": 6500000, "FINANCE": 7000000, "UPI": 3540000}


def test_type_and_finance_filters(sales):
    assert _report(sales, types=["SERVICE"]).sale_value == 35000
    assert _report(sales, types=["VEHICLE"], finance="LOAN").count == 1
    direct = _report(sales, types=["VEHICLE"], finance="DIRECT")
    assert direct.count == 1 and direct.by_mode == {"CASH": 5000000, "UPI": 3500000}
    assert _report(sales, types=[]).count == 0


def test_payment_mode_filter_counts_only_that_mode(sales):
    upi = _report(sales, modes=["UPI"])
    assert [i["invoice_type"] for i in upi.invoices] == ["VEHICLE", "PART"]
    assert upi.received_in_modes == 3500000 + 40000  # not the full 85,000 of the split vehicle bill
    assert upi.by_mode == {"UPI": 3540000}
    cash = _report(sales, modes=["CASH"])
    assert cash.received_in_modes == 6500000
    both = _report(sales, modes=["CASH", "UPI"])
    assert both.received_in_modes == 6500000 + 3540000  # split bill counted once per mode, never doubled


def test_all_dates(sales):
    assert sales_report(sales, Filters()).count == 5


def test_sales_report_page(qtbot, sales, tmp_path, monkeypatch, opened_files):
    from openpyxl import load_workbook

    from jbhonda.ui import sales_report_page
    from jbhonda.ui.sales_report_page import SalesReportPage

    page = SalesReportPage(sales)
    qtbot.addWidget(page)
    page.period.setValue("ALL")
    assert page.cards["count"].value.text() == "5"
    assert page.table.isColumnHidden(6)
    page.all_modes.setChecked(False)
    page.mode_boxes["UPI"].setChecked(True)
    assert page.cards["received"].value.text() == "₹35,400.00"
    assert not page.table.isColumnHidden(6)
    page.type_boxes["VEHICLE"].setChecked(False)
    assert page.cards["count"].value.text() == "1"

    target = tmp_path / "report.xlsx"
    monkeypatch.setattr(sales_report_page, "_choose_save_path", lambda parent, suggested: target)
    assert page.export() == target and opened_files == [target]
    assert load_workbook(target).sheetnames == ["Summary", "Invoices", "By payment mode"]
